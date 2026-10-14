"""list0722.xlsx + 덤프 보조 데이터(JSON) → Google 시트로 올릴 정리된 워크북 생성.

사용:
  python -I build_sheet.py <원본 xlsx> <보조 JSON 폴더> <출력 xlsx>

원칙
- 엑셀 37컬럼은 원문 그대로(순서 유지). 빈 열 200여 개는 버린다.
- 파란 셀(해외 보관) → '해외보관' 열 값 Y.
- 덤프에만 있던 값(시대·카테고리·지역·언어·이야기·선택 이미지)은 뒤쪽 열에 붙인다.
- 사람이 시트에서 바로 읽고 고칠 수 있게 코드값 대신 한글 표기를 쓴다.
"""

from __future__ import annotations

import datetime as dt
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC, AUX, OUT = (Path(a) for a in sys.argv[1:4])

ORIGINAL_HEADERS = [
    "No.", "MEMO-1", "Image", "Item No.", "TITLE", "Maker (CN)", "Type", "Year", "Size",
    "Condition", "신청", "신청 (전)", "신청 (전전)", "배송", "Location", "QTY", "정가 (원)",
    "판매가 (US$)", "MEMO-2", "MEMO-3", "한국원가", "현지구입가", "원가분석", "구매처", "결제일",
    "결제방법", "Lot#", "MEMO-4", "MEMO-5", "작품설명 (한글)", "작품설명 (Non-한글)", "Site",
    "MEMO-6", "MEMO-7", "MEMO-8", "MEMO-9", "MEMO-10",
]
EXTRA_HEADERS = [
    "해외보관", "연도from", "연도to", "시대", "카테고리", "지역", "언어",
    "Shop게시", "이미지폴더", "대표이미지", "이야기", "수정일시", "수정자",
]
HEADERS = ORIGINAL_HEADERS + EXTRA_HEADERS
HEADER_ALIASES = {"Item": "Item No.", "No": "No."}
SHEET_PREFIX_TYPE = {"AB": "BOOK", "AM": "MAP", "AO": "OTHER", "AP": "PRINT"}
VALID_TYPES = set(SHEET_PREFIX_TYPE.values())
OVERSEAS_FILL_RGB = "FF00B0F0"
_YEAR_RE = re.compile(r"(?<!\d)(1[0-9]{3}|20[0-9]{2})(?!\d)")


def cell_text(v):
    """원문 보존용 값. 날짜형은 ISO 문자열로, 숫자는 숫자 그대로, 공백뿐이면 None."""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    s = str(v).strip()
    return s or None


def year_range(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.year, None
    s = cell_text(v)
    if s is None:
        return None, None
    years = [int(m) for m in _YEAR_RE.findall(str(s))]
    if not years:
        return None, None
    lo, hi = min(years), max(years)
    return lo, (hi if hi > lo else None)


def fill_rgb(cell):
    f = cell.fill
    if f is not None and f.patternType == "solid" and f.fgColor is not None:
        if f.fgColor.type == "rgb" and isinstance(f.fgColor.rgb, str):
            return f.fgColor.rgb
    return None


def load_json(name):
    p = AUX / f"{name}.json"
    return json.loads(p.read_text()) if p.is_file() and p.stat().st_size else []


# ---- 보조 데이터 ------------------------------------------------------------
extra_by_code = {r["item_code"]: r for r in load_json("items_extra")}
cats_by_code: dict[str, list[str]] = defaultdict(list)
for r in load_json("item_categories"):
    name = r["category"]
    if name not in cats_by_code[r["item_code"]]:
        cats_by_code[r["item_code"]].append(name)
eras_by_code: dict[str, list[str]] = defaultdict(list)
for r in load_json("item_eras"):
    if r["relation_type"] not in ("production", "subject", "primary", "related"):
        pass
    label = r["era_name"] if r["is_primary"] else r["era_name"]
    if label not in eras_by_code[r["item_code"]]:
        eras_by_code[r["item_code"]].append(label)
stories = {r["item_code"]: r["story"] for r in load_json("item_stories")}
media = {}
for r in load_json("item_drive_media"):
    images = r["images"] if isinstance(r["images"], list) else json.loads(r["images"] or "[]")
    media[r["item_code"]] = (r["selected_folder_id"], images[0]["file_id"] if images else None)

# ---- 원본 파싱 -------------------------------------------------------------
wb_src = load_workbook(SRC, data_only=True)
rows_by_type: dict[str, list[list]] = {t: [] for t in VALID_TYPES}
seen: dict[str, str] = {}
report = defaultdict(int)
dups = []

for ws in wb_src.worksheets:
    prefix = ws.title[:2]
    default_type = SHEET_PREFIX_TYPE.get(prefix)
    if default_type is None:
        print("건너뜀 시트:", ws.title)
        continue
    header_idx = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(1, col).value
        if raw is None:
            continue
        h = HEADER_ALIASES.get(str(raw).strip(), str(raw).strip())
        header_idx.setdefault(h, col)
    missing = [h for h in ORIGINAL_HEADERS if h not in header_idx]
    if missing:
        print(f"[{ws.title}] 없는 헤더:", missing)
    code_col = header_idx["Item No."]
    for row_no in range(2, ws.max_row + 1):
        code = cell_text(ws.cell(row_no, code_col).value)
        if code is None:
            if any(ws.cell(row_no, c).value not in (None, "") for c in header_idx.values()):
                report[f"{ws.title}:코드없음"] += 1
            continue
        code = str(code)
        if code in seen:
            dups.append((ws.title, row_no, code, seen[code]))
            continue
        seen[code] = f"{ws.title}:{row_no}"
        values = [cell_text(ws.cell(row_no, header_idx[h]).value) if h in header_idx else None for h in ORIGINAL_HEADERS]
        rec = dict(zip(ORIGINAL_HEADERS, values))
        t = str(rec["Type"] or "").upper()
        if t not in VALID_TYPES:
            if rec["Type"]:
                report[f"{ws.title}:Type대체"] += 1
            t = default_type
        rec["Type"] = t
        location = rec["Location"]
        overseas = (str(location).strip().startswith("00") if location else False) or fill_rgb(ws.cell(row_no, code_col)) == OVERSEAS_FILL_RGB
        extra = extra_by_code.get(code, {})
        if extra.get("is_overseas"):
            overseas = True
        y_from, y_to = year_range(ws.cell(row_no, header_idx["Year"]).value)
        folder_id, file_id = media.get(code, (None, None))
        rec.update({
            "해외보관": "Y" if overseas else None,
            "연도from": y_from,
            "연도to": y_to,
            "시대": "; ".join(eras_by_code.get(code, [])) or extra.get("era"),
            "카테고리": "; ".join(cats_by_code.get(code, [])) or None,
            "지역": extra.get("region"),
            "언어": extra.get("language"),
            "Shop게시": None,
            "이미지폴더": folder_id,
            "대표이미지": file_id,
            "이야기": stories.get(code),
            "수정일시": None,
            "수정자": None,
        })
        rows_by_type[t].append([rec[h] for h in HEADERS])
        report[f"{ws.title}:적재"] += 1

# ---- 출력 워크북 ------------------------------------------------------------
wb = Workbook()
wb.remove(wb.active)
head_font = Font(bold=True)
head_fill = PatternFill("solid", fgColor="FFEDE7D8")
extra_fill = PatternFill("solid", fgColor="FFF7F1E3")
WIDTHS = {"TITLE": 44, "Maker (CN)": 28, "작품설명 (한글)": 50, "작품설명 (Non-한글)": 50, "이야기": 50,
          "Item No.": 18, "원가분석": 24, "카테고리": 24, "시대": 20}


def style_header(ws, headers, extra_from=None):
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i)
        c.font = head_font
        c.fill = extra_fill if (extra_from is not None and i > extra_from) else head_fill
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, max(10, min(18, len(h) + 4)))
    ws.freeze_panes = "E2" if "Item No." in headers else "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


for t in ["BOOK", "MAP", "OTHER", "PRINT"]:
    ws = wb.create_sheet(t)
    style_header(ws, HEADERS, extra_from=len(ORIGINAL_HEADERS))
    rows = sorted(rows_by_type[t], key=lambda r: str(r[3]))
    for r in rows:
        ws.append(r)
        if r[HEADERS.index("해외보관")] == "Y":
            ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=OVERSEAS_FILL_RGB)
    print(t, len(rows))

# 동의어
ws = wb.create_sheet("동의어")
style_header(ws, ["개념ID", "대표어", "유형", "표현", "관계", "가중치", "사용"])
for r in load_json("synonyms"):
    if r.get("term") is None:
        continue
    ws.append([r["concept_id"], r["canonical_term"], r.get("concept_type"), r["term"],
               r.get("relation_type"), r.get("search_weight"), "Y" if r.get("approved", True) else None])

# 카테고리
ws = wb.create_sheet("카테고리")
style_header(ws, ["대분류", "소분류", "순서", "사용"])
for r in load_json("categories"):
    top, sub = (r["name"], None) if r["parent"] is None else (r["parent"], r["name"])
    ws.append([top, sub, r.get("sort_order"), "Y" if r.get("is_active", True) else None])

# 시대
ws = wb.create_sheet("시대")
style_header(ws, ["코드", "이름", "시작연도", "끝연도", "지역", "설명", "순서", "사용"])
for r in load_json("eras"):
    ws.append([r["code"], r["name"], r.get("year_from"), r.get("year_to"), r.get("region"),
               r.get("description"), r.get("sort_order"), "Y" if r.get("is_active", True) else None])

# 이미지 폴더 (Drive 루트)
ws = wb.create_sheet("이미지폴더")
style_header(ws, ["이름", "폴더ID", "폴더URL", "사용", "순서"])
for r in load_json("drive_sources"):
    ws.append([r["name"], r["folder_id"], r["folder_url"], "Y" if r.get("is_active") else None, r.get("sort_order")])

# 설정/안내
ws = wb.create_sheet("설정")
style_header(ws, ["키", "값", "설명"])
ws.append(["schema_version", "1", "시트 구조 버전. 어드민이 확인한다"])
ws.append(["item_tabs", "BOOK,MAP,OTHER,PRINT", "물건 탭 이름"])
ws.append(["key_column", "Item No.", "물건 고유 키 열"])
ws.append(["source", SRC.name, "이관 원본 파일"])
ws.append(["migrated_at", dt.datetime.now().isoformat(timespec="seconds"), "이관 시각"])
ws.append(["note", "앞 37열은 엑셀 원문. 연한 바탕의 뒤쪽 열은 어드민이 쓰는 보조 열", ""])
ws.column_dimensions["B"].width = 40
ws.column_dimensions["C"].width = 50

wb.save(OUT)
print("---- 리포트")
for k in sorted(report):
    print(k, report[k])
print("중복 코드(첫 행만 유지):", len(dups))
for d in dups[:10]:
    print("  ", d)
print("총 물건:", sum(len(v) for v in rows_by_type.values()))
print("출력:", OUT)
