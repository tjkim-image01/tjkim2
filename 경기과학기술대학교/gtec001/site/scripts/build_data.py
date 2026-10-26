"""졸업전시 자료 → 사이트 데이터.
입력: 경기과학기술대학교/ 의 분야별 학생 폴더 · 졸업자 싸인 · 도록 PDF · 표지 PDF
출력: public/artists/<id>/*.jpg · public/sign/<id>.png · public/brand/* · public/catalog/* · src/data/artists.json · src/data/exhibition.json
실행: python -X utf8 -I scripts/build_data.py   (site 폴더에서)
"""
import os, re, json, colorsys
from pathlib import Path
from PIL import Image, ImageOps
import pymupdf

SITE = Path(__file__).resolve().parents[1]
RAW = SITE.parents[1]            # 경기과학기술대학교/
PUB = SITE / "public"; DATA = SITE / "src/data"
for d in ["artists", "sign", "brand", "catalog"]: (PUB / d).mkdir(parents=True, exist_ok=True)
DATA.mkdir(parents=True, exist_ok=True)

DEPTS = [("webtoon", "웹툰", "Webtoon", "#E98AA8"), ("illust", "일러스트", "Illustration", "#6FA8DC"),
         ("comics", "출판만화", "Comics", "#F2A15A"), ("anim", "애니메이션", "Animation", "#6CC0A0")]
DEPT_BY_EN = {en: code for code, ko, en, _ in DEPTS}

# ---------- 1. 도록 PDF 에서 학생 페이지 읽기 (편집된 최종 문안) ----------
cat_pdf = next(RAW.glob("*졸전도록.pdf"))
doc = pymupdf.open(cat_pdf)
from_pdf = {}
for pno, pg in enumerate(doc):
    txt = pg.get_text()
    if "학생 소개" not in txt: continue
    blocks = [(b[0], b[1], b[2], b[3], b[4].strip()) for b in pg.get_text("blocks") if b[4].strip()]
    dept = None; name = name_en = email = sns = title = ""; intro = []; note = []
    for x0, y0, x1, y1, t in blocks:
        if x0 < 70 and t in DEPT_BY_EN: dept = DEPT_BY_EN[t]; continue
        if x0 < 70: continue
        if t in ("학생 소개", "작품 제목", "작품소개", "소감, 짧은 한마디", "작품소개\n소감, 짧은 한마디"): continue
        if 95 < y0 < 122 and x0 > 180:
            lines = t.split("\n"); name = lines[0].strip(); name_en = " ".join(lines[1:]).strip()
        elif 122 < y0 < 142 and x0 > 180: email = t
        elif 142 < y0 < 165 and x0 > 180: sns = t
        elif 185 < y0 < 215 and x0 > 180: title = t.replace("\n", " ")
        elif y0 > 650 and x0 < 315: intro.append(t)
        elif y0 > 650 and x0 >= 315: note.append(t)
    join = lambda L: re.sub(r"\s*\n\s*", "", " ".join(L)).strip()
    if name:
        from_pdf[name] = dict(dept=dept, name=name, name_en=name_en, email=email, sns=sns, title=title,
                              intro=join(intro), note=join(note), page=pno + 1)
print("pdf students:", len(from_pdf))

# ---------- 1b. 학생이 낸 원문(txt·docx·hwp·hwpx) — 띄어쓰기가 정확하므로 소개·소감은 원문을 우선 ----------
import zipfile, zlib, struct
def read_txt(p):
    b = open(p, "rb").read()
    for enc in ("utf-8-sig", "cp949", "utf-16"):
        try: return b.decode(enc)
        except Exception: pass
    return b.decode("utf-8", "replace")
def read_docx(p):
    import docx
    d = docx.Document(p); out = [x.text for x in d.paragraphs if x.text.strip()]
    for t in d.tables:
        for r in t.rows: out.append(" | ".join(c.text for c in r.cells))
    return "\n".join(out)
def read_hwpx(p):
    z = zipfile.ZipFile(p); out = []
    for n in z.namelist():
        if n.startswith("Contents/section"):
            x = z.read(n).decode("utf-8"); out += [t for t in re.findall(r"<hp:t[^>]*>([^<]*)</hp:t>", x) if t.strip()]
    return "\n".join(out)
def read_hwp(p):
    import olefile
    ole = olefile.OleFileIO(p); hdr = ole.openstream("FileHeader").read()
    compressed = struct.unpack("<I", hdr[36:40])[0] & 1; out = []
    for e in ole.listdir():
        if e[0] == "BodyText":
            data = ole.openstream(e).read()
            if compressed: data = zlib.decompress(data, -15)
            i = 0
            while i + 4 <= len(data):
                h = struct.unpack("<I", data[i:i + 4])[0]; tag = h & 0x3ff; size = (h >> 20) & 0xfff; i += 4
                if size == 0xfff: size = struct.unpack("<I", data[i:i + 4])[0]; i += 4
                if tag == 67:
                    s = data[i:i + size].decode("utf-16le", "replace")
                    s = "".join(c if ord(c) >= 32 else "\n" if ord(c) in (13, 10) else "" for c in s)
                    if s.strip(): out.append(s.strip())
                i += size
    return "\n".join(out)
READERS = {".txt": read_txt, ".docx": read_docx, ".hwpx": read_hwpx, ".hwp": read_hwp}
LABEL = re.compile(r"^\s*(?:이름\s*,?\s*영문\s*이름|메일\s*주소|작품\s*제목|작품\s*소개|소감\s*,?\s*짧은\s*한\s*마디|소감)?\s*(?:\(최대\s*\d+자\))?\s*(?:\([^)]*통일해주세요\.?\))?\s*[:：\-–]?\s*")
def clean(s): return re.sub(r"[ \t]+", " ", LABEL.sub("", s or "").strip()).strip()
def parse_fields(t):
    t = t.replace("\r", "").replace("捤獥汤捯", "").replace("위의 내용을 복사해서 써주세요!", "")
    t = re.sub(r"^\s*도록용 글\s*", "", t); t = re.sub(r"(?m)^탭 \d\s*$", "", t)
    parts = re.split(r"(?:^|\n)\s*([1-6])\s*(?:[.)]\s*|\s+)(?=\S)", "\n" + t)
    segs = {}
    for i in range(1, len(parts) - 1, 2): segs.setdefault(int(parts[i]), parts[i + 1].strip())
    if len(segs) >= 5:
        if 6 in segs: return clean(segs.get(5)), clean(segs.get(6))
        return clean(segs.get(4)), clean(segs.get(5))          # SNS 항목을 건너뛴 경우
    mi = re.search(r"작품\s*소개\s*[:：]?\s*(.*?)(?=\n\s*소감|\Z)", t, re.S)
    mn = re.search(r"소감[,，]?\s*(?:짧은\s*한\s*마디)?\s*[:：]?\s*(.*)", t, re.S)
    return clean(mi.group(1)) if mi else "", clean(mn.group(1)) if mn else ""
def raw_texts(folder):
    """학생 폴더의 도록용 글 파일 → (소개, 소감). 졸준위용 파일은 뺀다."""
    for f in sorted(os.listdir(folder)):
        ext = os.path.splitext(f)[1].lower()
        if ext in READERS and "졸준위" not in f and "총무" not in f:
            try:
                intro, note = parse_fields(READERS[ext](os.path.join(folder, f)))
                if intro or note: return intro, note, f
            except Exception as ex: print("  raw read fail", f, ex)
    return "", "", None
OVR = json.load(open(SITE / "scripts/overrides.json", encoding="utf-8")) if (SITE / "scripts/overrides.json").exists() else {}

# ---------- 2. 폴더 → 이미지 ----------
def num_key(f):
    m = re.match(r"^\D*?(\d+)(?:-(\d+))?", f)
    return (int(m.group(1)), int(m.group(2) or 0)) if m else (99, 0)

def flatten(im):
    im = ImageOps.exif_transpose(im)
    if im.mode in ("P", "LA", "RGBA"):
        im = im.convert("RGBA"); bg = Image.new("RGB", im.size, (255, 255, 255)); bg.paste(im, mask=im.split()[3]); return bg
    return im.convert("RGB")

def save_jpg(im, path, max_side=None, max_w=None, q=84):
    im = flatten(im)
    w, h = im.size; s = 1.0
    if max_side: s = min(s, max_side / max(w, h))
    if max_w: s = min(s, max_w / w)
    if s < 1: im = im.resize((max(1, round(w * s)), max(1, round(h * s))), Image.LANCZOS)
    im.save(path, "JPEG", quality=q, optimize=True, progressive=True)
    return im.size

def palette(im):
    sm = flatten(im).resize((80, 80))
    q = sm.quantize(colors=8, method=Image.Quantize.MEDIANCUT).convert("RGB")
    counts = sorted(q.getcolors(80 * 80), reverse=True)
    cols = [c for _, c in counts]
    def hsv(c): return colorsys.rgb_to_hsv(*[v / 255 for v in c])
    base = cols[0]
    sat = [c for c in cols if hsv(c)[1] > 0.28 and 0.18 < hsv(c)[2] < 0.95]
    accent = max(sat, key=lambda c: hsv(c)[1] * (0.5 + hsv(c)[2])) if sat else cols[min(1, len(cols) - 1)]
    dark = min(cols, key=lambda c: hsv(c)[2])
    hexs = lambda c: "#%02X%02X%02X" % c
    mix = lambda c, w, t: tuple(round(c[i] * (1 - t) + w[i] * t) for i in range(3))
    return dict(base=hexs(base), accent=hexs(accent), dark=hexs(dark),
                wall=hexs(mix(base, (250, 247, 240), 0.72)), floor=hexs(mix(accent, (40, 36, 34), 0.55)),
                mood="dark" if hsv(base)[2] < 0.45 else "light")

SIGN_ROTATE = {"안소희": 90}  # 2026-10-24 사용자 요청: 안소희 서명 시계 반대 방향 90°
MAIN_PICK = {"김지완": "2.png"}  # 2026-10-25 사용자 요청: 대표 그림을 1-2 대신 이 파일로(1-2 는 작품 p1 로 감)
sign_dir = next(RAW.glob("졸업자 싸인*")) / "졸업자 싸인"
sign_files = {p.name: p for p in sign_dir.glob("*.png")}

artists = []
for code, ko, en, color in DEPTS:
    dept_dir = next(RAW.glob(f"{ko}-2026*")) / ko
    for st in sorted(os.listdir(dept_dir)):
        sp = dept_dir / st
        if not sp.is_dir() or "공지" in st: continue
        m = re.match(r"(\d+)[_ ]+(.+)", st); sid, name = m.group(1), m.group(2).strip()
        src = sp
        subs = [d for d in sp.iterdir() if d.is_dir() and "인스타" not in d.name]
        if subs: src = subs[0]
        files = [f for f in os.listdir(src) if f.lower().endswith((".png", ".jpg", ".jpeg")) and "업로드" not in f and "싸인" not in f and "서명" not in f]
        files.sort(key=num_key)
        info = from_pdf.get(name) or {}
        r_intro, r_note, r_file = raw_texts(src)
        aid = f"{code[:2]}{len([a for a in artists if a['dept'] == code]) + 1:02d}"
        out = PUB / "artists" / aid; out.mkdir(exist_ok=True)
        rec = dict(id=aid, dept=code, dept_ko=ko, dept_en=en, student_no=sid, name=name,
                   name_en=info.get("name_en", ""), email=info.get("email", ""), sns=info.get("sns", ""),
                   title=info.get("title", ""), intro=r_intro or info.get("intro", ""), note=r_note or info.get("note", ""),
                   intro_src="원문 " + r_file if r_intro else "도록 PDF", catalog_page=info.get("page"),
                   avatar=None, main=None, works=[], sign=None, palette=None)
        rec.update(OVR.get(name, {}))
        works = []
        for f in files:
            k = num_key(f); im = Image.open(src / f)
            if k == (1, 1):
                save_jpg(im, out / "avatar.jpg", max_side=320); rec["avatar"] = "avatar.jpg"
            elif (f == MAIN_PICK[name]) if name in MAIN_PICK else k == (1, 2):
                w, h = save_jpg(im, out / "main.jpg", max_side=1600)
                save_jpg(im, out / "main_thumb.jpg", max_side=640, q=80)
                save_jpg(im, out / "main_tex.jpg", max_side=1024, q=78)
                rec["main"] = dict(file="main.jpg", thumb="main_thumb.jpg", tex="main_tex.jpg", w=w, h=h)
                rec["palette"] = palette(im)
            else:
                works.append((k, f, im))
        for i, (k, f, im) in enumerate(works, 1):
            w0, h0 = im.size; strip = h0 / w0 > 2.2
            w, h = save_jpg(im, out / f"p{i}.jpg", max_side=None if strip else 1800, max_w=1000 if strip else None)
            save_jpg(im, out / f"p{i}_thumb.jpg", max_side=900 if strip else 640, q=80)
            save_jpg(im, out / f"p{i}_tex.jpg", max_side=1400 if strip else 1024, q=78)
            rec["works"].append(dict(file=f"p{i}.jpg", thumb=f"p{i}_thumb.jpg", tex=f"p{i}_tex.jpg", w=w, h=h, strip=strip, src=f))
        if rec["main"] is None and rec["works"]:
            first = rec["works"][0]; rec["main"] = first; rec["palette"] = palette(Image.open(out / first["file"]))
        # 서명
        cand = [p for n, p in sign_files.items() if name in n]
        for alt in ("싸인.png", "서명.png"):
            if not cand and (src / alt).exists(): cand = [src / alt]
        if cand:
            s = Image.open(cand[0]).convert("RGBA")
            px = s.getpixel((0, 0))
            if px[3] > 0 and min(px[:3]) > 235:
                s.putdata([(r, g, b, 0) if min(r, g, b) > 225 else (r, g, b, a) for r, g, b, a in s.getdata()])
            bbox = s.getbbox()
            if bbox: s = s.crop(bbox)
            if name in SIGN_ROTATE: s = s.rotate(SIGN_ROTATE[name], expand=True)  # 세로로 제출된 서명 바로 세우기(반시계 방향 각도)
            s.thumbnail((700, 700)); s.save(PUB / "sign" / f"{aid}.png"); rec["sign"] = f"{aid}.png"
        artists.append(rec)
        print(f"{aid} {ko:5s} {name:4s} {rec['title'][:18]:20s} imgs={len(files)} sign={'O' if rec['sign'] else '-'} pdf={'O' if info else '-'}")

missing_pdf = [a["name"] for a in artists if not a["catalog_page"]]
print("no pdf match:", missing_pdf, "| pdf-only:", set(from_pdf) - {a["name"] for a in artists})

# ---------- 3. 전시 공통 자료 ----------
cover_pdf = next(RAW.glob("*표지.pdf")); cdoc = pymupdf.open(cover_pdf)
pix = cdoc[2].get_pixmap(dpi=170); tmp = PUB / "brand" / "_cover.png"; pix.save(tmp)
Image.open(tmp).convert("RGB").save(PUB / "brand" / "cover.jpg", quality=86); tmp.unlink()
pix = doc[0].get_pixmap(dpi=200, alpha=False); tmp = PUB / "brand" / "_logo_page.png"  # 로고 = 도록 1쪽; pix.save(tmp)
lp = Image.open(tmp).convert("RGB")
W, H = lp.size; top = lp.crop((0, 0, W, int(H * 0.5)))
g = ImageOps.invert(ImageOps.grayscale(top)).point(lambda v: 255 if v > 40 else 0); bb = g.getbbox()
logo = top.crop((max(0, bb[0] - 30), max(0, bb[1] - 30), bb[2] + 30, bb[3] + 30)); logo.save(PUB / "brand" / "logo_paper.jpg", quality=90)
tmp.unlink()
cat_pages = {}
for p in [6, 7, 8, 9, 10, 11, 12, 13, 191, 192, 193, 194, 196, 198]:
    tmp = PUB / "catalog" / f"p{p:03d}.png"; doc[p - 1].get_pixmap(dpi=100).save(tmp)
    Image.open(tmp).convert("RGB").save(PUB / "catalog" / f"p{p:03d}.jpg", quality=84); tmp.unlink(); cat_pages[p] = f"p{p:03d}.jpg"
def page_text(p): return doc[p - 1].get_text().strip()
greet = page_text(6)
greet_body = re.sub(r"^\d+\s*", "", greet).split("고원상 학과장님 인사말")[0].strip()
greet_lines = [l.strip() for l in greet_body.split("\n") if l.strip()]
title_line = greet_lines[0]; paras = []; buf = ""
for l in greet_lines[1:]:
    buf = (buf + l) if (buf and not re.search(r"[.!?]$", buf)) else ((buf + " " + l).strip() if buf else l)
    if re.search(r"[.!?]$", buf) and len(buf) > 60: paras.append(buf.strip()); buf = ""
if buf: paras.append(buf.strip())
committee_raw = page_text(196)
committee = []
roles = r"(졸업전시준비위원장|총무부장|홍보부장|디피부장|도록부장)"
parts = re.split(roles + r"\s+(\S+)\s*\n", committee_raw)
for i in range(1, len(parts) - 2, 3):
    committee.append(dict(role=parts[i], name=parts[i + 1], text=re.sub(r"\s*\n\s*", "", parts[i + 2]).strip()))
# 졸준위 글은 학생 폴더의 원문(띄어쓰기 정확)이 있으면 그걸 쓴다
for c in committee:
    for a in artists:
        if a["name"] != c["name"]: continue
        folder = next((d for d in (RAW.glob(f"{a['dept_ko']}-2026*/{a['dept_ko']}/{a['student_no']}*")) ), None)
        if not folder: continue
        for dp, dn, fn in os.walk(folder):
            for f in fn:
                if ("졸준위" in f or "총무" in f) and os.path.splitext(f)[1].lower() in READERS:
                    try: c["text"] = re.sub(r"\s+", " ", READERS[os.path.splitext(f)[1].lower()](os.path.join(dp, f)).replace("捤獥汤捯", "")).strip(); c["src"] = f
                    except Exception as ex: print("committee raw fail", f, ex)
    if c["name"] in OVR.get("_committee", {}): c["text"] = OVR["_committee"][c["name"]]
if "_greeting" in OVR: title_line, paras = OVR["_greeting"]["title"], OVR["_greeting"]["paras"]
colophon = re.sub(r"\s*\n\s*", " ", page_text(197))
ex = dict(
    school="경기과학기술대학교", dept="웹툰일러스트학과", edition="제4회 졸업전시회", title="MOVE : 이소",
    subtitle="여정의 끝, 그리고 또다른 시작", en="webtoon&Illustration", dates="2026. 10. 28 – 11. 02",
    date_start="2026-10-28", date_end="2026-11-02", counts={c: len([a for a in artists if a["dept"] == c]) for c, *_ in DEPTS},
    depts=[dict(code=c, ko=k, en=e, color=col) for c, k, e, col in DEPTS],
    greeting=dict(title=title_line, paras=paras, by="웹툰일러스트학과 학과장 고원상"),
    tributes=[dict(page=7, who="고원상 교수님 축전", img=cat_pages[7]), dict(page=8, who="원민관 교수님 축전", img=cat_pages[8]),
              dict(page=9, who="박미승 교수님 축전", img=cat_pages[9]), dict(page=10, who="고홍석 교수님 축전", img=cat_pages[10]),
              dict(page=11, who="김대현 교수님 축전", img=cat_pages[11]), dict(page=12, who="황기연 교수님 축전", img=cat_pages[12]),
              dict(page=13, who="임성구 예스웍스 대표님 축전", img=cat_pages[13]), dict(page=192, who="김나윤 선배 축전", img=cat_pages[192]),
              dict(page=193, who="임성택 선배 축전", img=cat_pages[193]), dict(page=194, who="최지혁 후배 축전", img=cat_pages[194])],
    committee=committee, colophon=colophon, signatures_page=cat_pages[198], cover="brand/cover.jpg", logo="brand/logo_paper.jpg",
)
# 도록에 없는 추가 축전(황지영 교수님, 2026-10-24 요청) — 이미지는 scripts/make_tribute_hwang.py 가 아니라 public/catalog/tribute_hwang.jpg 로 이미 만들어 둠
if (PUB / "catalog" / "tribute_hwang.jpg").exists():
    ex["tributes"].insert(7, dict(page=None, who="황지영 교수님 축전", img="tribute_hwang.jpg", src="퓨어데이 가족 포스터 v2(이마리·조성민·조하임) + 축전 문안"))
# 도록 뒤에 받은 축전(2026-10-25 요청) — 원본 PNG(경기과학기술대학교/KakaoTalk_20261025_*.png)를 1240×1754 jpg 로 줄여 public/catalog 에 둠. 선배 축전 뒤, 후배 축전 앞
_late = [dict(page=None, who="박상빈 선배 축전", img="tribute_parksangbin.jpg", src="KakaoTalk_20261025_201148773.png, 1기 졸업생"),
         dict(page=None, who="이예원 축전 (나서연 지인)", img="tribute_leeyewon.jpg", src="KakaoTalk_20261025_201203067.png")]
_at = next((i for i, t in enumerate(ex["tributes"]) if "후배" in t["who"]), len(ex["tributes"]))
ex["tributes"][_at:_at] = [t for t in _late if (PUB / "catalog" / t["img"]).exists()]
# 도록 뒤에 받은 서명 199쪽(2026-10-26 요청) — 원본(수정사항/KakaoTalk_20261026_100227821.jpg)을 p198 과 같은 827×1170 으로 줄여 둠. 198쪽 옆에 나란히
ex["signatures_pages"] = [cat_pages[198]] + [f for f in ["sign_p199.jpg"] if (PUB / "catalog" / f).exists()]
json.dump(artists, open(DATA / "artists.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
json.dump(ex, open(DATA / "exhibition.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("greeting paras:", len(paras), "| committee:", [(c["role"], c["name"]) for c in committee])
print("done", len(artists))
